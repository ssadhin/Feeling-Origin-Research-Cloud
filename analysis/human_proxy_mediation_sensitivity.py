from __future__ import annotations

import argparse
import json
import math
import urllib.request
from collections import defaultdict
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

DATA_COMMIT = "692d57f47861684b5144cfd4f3e6f0c9a2c3dc43"
DATA_URL = (
    "https://raw.githubusercontent.com/EP171993/painperception/"
    + DATA_COMMIT
    + "/Data_v3%20-%202504.xlsx"
)
MEDIATORS = ("HR", "HRd", "HRd_Z")
BASE_REQUIRED = {"Exp", "Feedback", "Trial", "SSID", "LIK", "VAS", "StimInt"}


def num(x):
    if x is None or x == "":
        return None
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def load_rows(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    iterator = ws.iter_rows(values_only=True)
    header = [str(v).strip() if v is not None else "" for v in next(iterator)]
    needed = BASE_REQUIRED | set(MEDIATORS)
    missing = needed - set(header)
    if missing:
        raise RuntimeError(f"Missing columns {sorted(missing)}; columns={header}")
    idx = {k: header.index(k) for k in needed}
    rows = []
    for vals in iterator:
        r = {k: vals[i] for k, i in idx.items()}
        if r["SSID"] is None:
            continue
        r["SSID"] = str(r["SSID"]).strip()
        r["Exp"] = str(r["Exp"]).strip()
        r["Feedback"] = str(r["Feedback"]).strip()
        for k in ("Trial", "StimInt", "LIK", "VAS", *MEDIATORS):
            r[k] = num(r[k])
        rows.append(r)
    return rows, header


def design(rows, predictors):
    subs = sorted({r["_sid"] for r in rows})
    if len(subs) < 2:
        raise RuntimeError("Need >=2 subjects")
    cols = ["intercept", *predictors, *[f"subj:{s}" for s in subs[1:]]]
    X = np.zeros((len(rows), len(cols)), float)
    X[:, 0] = 1
    for j, p in enumerate(predictors, 1):
        X[:, j] = [float(r[p]) for r in rows]
    base = 1 + len(predictors)
    lookup = {s: base + j for j, s in enumerate(subs[1:])}
    for i, r in enumerate(rows):
        if r["_sid"] in lookup:
            X[i, lookup[r["_sid"]]] = 1
    return X, cols


def fit(rows, outcome, predictors):
    X, cols = design(rows, predictors)
    y = np.asarray([float(r[outcome]) for r in rows])
    beta, _, rank, _ = np.linalg.lstsq(X, y, rcond=None)
    resid = y - X @ beta
    sse = float(resid @ resid)
    sst = float(((y-y.mean())**2).sum())
    return {
        "coef": {c: float(b) for c, b in zip(cols, beta)},
        "rank": int(rank),
        "r2": None if sst <= 1e-15 else float(1-sse/sst),
    }


def prep(rows, exp, outcome, mediator):
    rr = []
    for r in rows:
        if r["Exp"] != exp or r["Feedback"] not in ("Faster", "Slower"):
            continue
        if r["StimInt"] in (None, 1.0, 5.0):
            continue
        if any(r[k] is None for k in ("Trial", "StimInt", outcome, mediator)):
            continue
        q = dict(r)
        q["_sid"] = r["SSID"]
        q["proxy"] = float(r["Feedback"] == "Faster")
        q["stim"] = (float(r["StimInt"]) - 3.0) / 2.0
        rr.append(q)
    if not rr:
        raise RuntimeError(f"No data {exp}/{outcome}/{mediator}")
    t = np.asarray([r["Trial"] for r in rr], float)
    m, s = float(t.mean()), float(t.std())
    s = s if s > 0 else 1.0
    for r in rr:
        r["trial_z"] = (r["Trial"] - m) / s
        r["proxy_x_trial"] = r["proxy"] * r["trial_z"]
    return rr


def estimate(rows, outcome, mediator, interaction=False):
    cov = ["stim", "trial_z"]
    treat_terms = ["proxy", "proxy_x_trial"] if interaction else ["proxy"]
    med_model = fit(rows, mediator, [*treat_terms, *cov])
    total = fit(rows, outcome, [*treat_terms, *cov])
    direct = fit(rows, outcome, [*treat_terms, mediator, *cov])
    a = med_model["coef"]["proxy"]
    b = direct["coef"][mediator]
    c = total["coef"]["proxy"]
    cp = direct["coef"]["proxy"]
    ind = a*b
    return {
        "a_proxy_to_mediator": a,
        "b_mediator_to_outcome_adjusted": b,
        "c_total_proxy_to_outcome": c,
        "c_prime_direct_proxy_to_outcome": cp,
        "indirect_a_times_b": ind,
        "mediated_fraction": None if abs(c) < 1e-12 else ind/c,
        "proxy_x_trial_mediator": med_model["coef"].get("proxy_x_trial"),
        "proxy_x_trial_total": total["coef"].get("proxy_x_trial"),
        "proxy_x_trial_direct": direct["coef"].get("proxy_x_trial"),
        "n_rows": len(rows),
        "n_subjects": len({r["_sid"] for r in rows}),
    }


def cluster_boot(rows, outcome, mediator, n, seed, interaction=False):
    rng = np.random.default_rng(seed)
    grouped = defaultdict(list)
    for r in rows:
        grouped[r["_sid"]].append(r)
    subs = sorted(grouped)
    keys = [
        "a_proxy_to_mediator", "b_mediator_to_outcome_adjusted",
        "c_total_proxy_to_outcome", "c_prime_direct_proxy_to_outcome",
        "indirect_a_times_b", "mediated_fraction",
        "proxy_x_trial_mediator", "proxy_x_trial_total", "proxy_x_trial_direct",
    ]
    vals = {k: [] for k in keys}
    fail = 0
    for _ in range(n):
        sample = rng.choice(subs, len(subs), replace=True)
        br = []
        for j, sid in enumerate(sample):
            for r in grouped[sid]:
                q = dict(r)
                q["_sid"] = f"{sid}#b{j}"
                br.append(q)
        try:
            e = estimate(br, outcome, mediator, interaction)
            for k in keys:
                v = e.get(k)
                if v is not None and math.isfinite(v):
                    vals[k].append(v)
        except Exception:
            fail += 1
    out = {}
    for k, vs in vals.items():
        if len(vs) < max(50, n//4):
            out[k] = None
        else:
            x = np.asarray(vs)
            out[k] = {
                "median": float(np.median(x)),
                "ci95": [float(np.percentile(x,2.5)), float(np.percentile(x,97.5))],
                "n_valid": len(vs),
            }
    return out, fail


def excludes_zero(item):
    if not item:
        return False
    lo, hi = item["ci95"]
    return lo > 0 or hi < 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--bootstrap", type=int, default=1000)
    ap.add_argument("--seed", type=int, default=15115)
    args = ap.parse_args()
    outdir = Path("results"); outdir.mkdir(exist_ok=True)
    path = outdir/"external_parrotta_data.xlsx"
    if not path.exists():
        urllib.request.urlretrieve(DATA_URL, path)
    rows, header = load_rows(path)
    result = {
        "status":"EXPLORATORY_SENSITIVITY_NOT_PREREGISTERED",
        "source":{
            "doi":"10.7554/eLife.90013",
            "repo":"EP171993/painperception",
            "dataset_commit":DATA_COMMIT,
            "dataset_blob_sha":"447fff57c91fb5ff761449d98d93737e26b1236c",
        },
        "firewall":(
            "Sensitivity analysis across HR, HRd and HRd_Z. These variables are ECG-derived summaries, "
            "not exhaustive bodily/evaluative sources. A non-mediated residual is not proof of a "
            "source-independent phenomenal effect; mediation assumptions and timing may fail."
        ),
        "analyses":{},
    }
    run_index=0
    for exp in ("Intero","Extero"):
        result["analyses"][exp]={}
        for outcome in ("LIK","VAS"):
            result["analyses"][exp][outcome]={}
            for mediator in MEDIATORS:
                rr=prep(rows,exp,outcome,mediator)
                p=estimate(rr,outcome,mediator,interaction=False)
                ci,fail=cluster_boot(rr,outcome,mediator,args.bootstrap,args.seed+run_index,False)
                # secondary sensitivity: allow treatment effect to vary over trial.
                pi=estimate(rr,outcome,mediator,interaction=True)
                cii,faili=cluster_boot(rr,outcome,mediator,args.bootstrap,args.seed+1000+run_index,True)
                result["analyses"][exp][outcome][mediator]={
                    "average_effect_model":{"point":p,"bootstrap":ci,"failures":fail},
                    "feedback_by_trial_model":{"point":pi,"bootstrap":cii,"failures":faili},
                }
                run_index += 1

    gate={}
    for outcome in ("LIK","VAS"):
        gate[outcome]={}
        for mediator in MEDIATORS:
            x=result["analyses"]["Intero"][outcome][mediator]["average_effect_model"]
            gate[outcome][mediator]={
                "proxy_changes_mediator":excludes_zero(x["bootstrap"]["a_proxy_to_mediator"]),
                "total_proxy_effect":excludes_zero(x["bootstrap"]["c_total_proxy_to_outcome"]),
                "indirect_detected":excludes_zero(x["bootstrap"]["indirect_a_times_b"]),
                "residual_direct_detected":excludes_zero(x["bootstrap"]["c_prime_direct_proxy_to_outcome"]),
                "indirect_ci":x["bootstrap"]["indirect_a_times_b"]["ci95"] if x["bootstrap"]["indirect_a_times_b"] else None,
                "direct_ci":x["bootstrap"]["c_prime_direct_proxy_to_outcome"]["ci95"] if x["bootstrap"]["c_prime_direct_proxy_to_outcome"] else None,
            }
    result["intero_sensitivity_gate"]=gate
    out=(outdir/"human_proxy_mediation_sensitivity.json")
    out.write_text(json.dumps(result,indent=2),encoding="utf-8")
    print(json.dumps(gate,indent=2))


if __name__=="__main__":
    main()
