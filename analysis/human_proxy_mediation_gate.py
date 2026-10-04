from __future__ import annotations

import argparse
import json
import math
import os
import urllib.request
from collections import defaultdict
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

DATA_COMMIT = "692d57f47861684b5144cfd4f3e6f0c9a2c3dc43"
DATA_URL = (
    "https://raw.githubusercontent.com/EP171993/painperception/"
    + DATA_COMMIT
    + "/Data_v3%20-%202504.xlsx"
)

REQUIRED = {"Exp", "Feedback", "Trial", "SSID", "HR", "LIK", "VAS", "StimInt"}


def as_float(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def load_rows(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = [str(x).strip() if x is not None else "" for x in next(it)]
    missing = REQUIRED - set(header)
    if missing:
        raise RuntimeError(f"Missing required columns: {sorted(missing)}; columns={header}")
    idx = {name: header.index(name) for name in REQUIRED}
    rows = []
    for vals in it:
        row = {name: vals[i] for name, i in idx.items()}
        if row["SSID"] is None:
            continue
        row["Exp"] = str(row["Exp"]).strip()
        row["Feedback"] = str(row["Feedback"]).strip()
        row["SSID"] = str(row["SSID"]).strip()
        for k in ("Trial", "HR", "LIK", "VAS", "StimInt"):
            row[k] = as_float(row[k])
        rows.append(row)
    return rows, header


def make_design(rows, predictor_names, subject_key="_sid"):
    subjects = sorted({r[subject_key] for r in rows})
    if len(subjects) < 2:
        raise RuntimeError("Need at least two subjects")
    subject_cols = subjects[1:]
    X = np.zeros((len(rows), 1 + len(predictor_names) + len(subject_cols)), dtype=float)
    X[:, 0] = 1.0
    names = ["intercept"] + predictor_names + [f"subj:{s}" for s in subject_cols]
    for j, name in enumerate(predictor_names, start=1):
        X[:, j] = [float(r[name]) for r in rows]
    base = 1 + len(predictor_names)
    subject_to_col = {s: base + j for j, s in enumerate(subject_cols)}
    for i, r in enumerate(rows):
        c = subject_to_col.get(r[subject_key])
        if c is not None:
            X[i, c] = 1.0
    return X, names


def ols(rows, outcome, predictors, subject_key="_sid"):
    X, names = make_design(rows, predictors, subject_key)
    y = np.asarray([float(r[outcome]) for r in rows], dtype=float)
    beta, _, rank, _ = np.linalg.lstsq(X, y, rcond=None)
    pred = X @ beta
    resid = y - pred
    sse = float(resid @ resid)
    sst = float(((y - y.mean()) ** 2).sum())
    return {
        "coef": {n: float(v) for n, v in zip(names, beta)},
        "rank": int(rank),
        "n": int(len(y)),
        "r2": None if sst <= 1e-15 else float(1.0 - sse / sst),
    }


def prepare(rows, exp, outcome):
    selected = []
    for r in rows:
        if r["Exp"] != exp or r["Feedback"] not in ("Faster", "Slower"):
            continue
        if r["StimInt"] in (None, 1.0, 5.0):
            continue
        if any(r[k] is None for k in ("Trial", "HR", outcome)):
            continue
        x = dict(r)
        x["_sid"] = r["SSID"]
        x["proxy"] = 1.0 if r["Feedback"] == "Faster" else 0.0
        x["stim"] = (float(r["StimInt"]) - 3.0) / 2.0
        selected.append(x)
    if not selected:
        raise RuntimeError(f"No analyzable rows for {exp}/{outcome}")
    trials = np.asarray([r["Trial"] for r in selected], dtype=float)
    sd = float(trials.std())
    mean = float(trials.mean())
    for r in selected:
        r["trial_z"] = (r["Trial"] - mean) / (sd if sd > 0 else 1.0)
    return selected


def paired_contrast(rows, outcome):
    by = defaultdict(lambda: {0: [], 1: []})
    for r in rows:
        by[r["_sid"]][int(r["proxy"])].append(float(r[outcome]))
    diffs = []
    for _, g in by.items():
        if g[0] and g[1]:
            diffs.append(float(np.mean(g[1]) - np.mean(g[0])))
    return {
        "n_subjects_with_both": len(diffs),
        "mean_faster_minus_slower": None if not diffs else float(np.mean(diffs)),
        "sd_subject_differences": None if len(diffs) < 2 else float(np.std(diffs, ddof=1)),
    }


def point_estimate(rows, outcome):
    # All three models use exactly the same rows.
    med = ols(rows, "HR", ["proxy", "stim", "trial_z"])
    total = ols(rows, outcome, ["proxy", "stim", "trial_z"])
    direct = ols(rows, outcome, ["proxy", "HR", "stim", "trial_z"])
    a = med["coef"]["proxy"]
    c = total["coef"]["proxy"]
    cp = direct["coef"]["proxy"]
    b = direct["coef"]["HR"]
    indirect = a * b
    frac = None if abs(c) < 1e-12 else indirect / c
    return {
        "a_proxy_to_HR": a,
        "b_HR_to_outcome_adjusted": b,
        "c_total_proxy_to_outcome": c,
        "c_prime_direct_proxy_to_outcome": cp,
        "indirect_a_times_b": indirect,
        "mediated_fraction": frac,
        "r2_mediator": med["r2"],
        "r2_total": total["r2"],
        "r2_direct": direct["r2"],
        "n_rows": len(rows),
        "n_subjects": len({r["_sid"] for r in rows}),
    }


def bootstrap(rows, outcome, n_boot, seed):
    rng = np.random.default_rng(seed)
    by = defaultdict(list)
    for r in rows:
        by[r["_sid"]].append(r)
    subjects = sorted(by)
    keys = [
        "a_proxy_to_HR",
        "b_HR_to_outcome_adjusted",
        "c_total_proxy_to_outcome",
        "c_prime_direct_proxy_to_outcome",
        "indirect_a_times_b",
        "mediated_fraction",
    ]
    vals = {k: [] for k in keys}
    failures = 0
    for _ in range(n_boot):
        sampled = rng.choice(subjects, size=len(subjects), replace=True)
        br = []
        for j, sid in enumerate(sampled):
            pseudo = f"{sid}#boot{j}"
            for r in by[sid]:
                q = dict(r)
                q["_sid"] = pseudo
                br.append(q)
        try:
            est = point_estimate(br, outcome)
            for k in keys:
                v = est[k]
                if v is not None and math.isfinite(v):
                    vals[k].append(float(v))
        except Exception:
            failures += 1
    cis = {}
    for k, arr in vals.items():
        if len(arr) < max(50, n_boot // 4):
            cis[k] = None
        else:
            a = np.asarray(arr)
            cis[k] = {
                "median": float(np.median(a)),
                "ci95": [float(np.percentile(a, 2.5)), float(np.percentile(a, 97.5))],
                "n_valid": int(len(a)),
            }
    return cis, failures


def analyze(rows, exp, outcome, n_boot, seed):
    rr = prepare(rows, exp, outcome)
    point = point_estimate(rr, outcome)
    cis, failures = bootstrap(rr, outcome, n_boot, seed)
    return {
        "point": point,
        "cluster_bootstrap_by_subject": cis,
        "bootstrap_failures": failures,
        "paired_raw_condition_contrast": paired_contrast(rr, outcome),
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--bootstrap", type=int, default=1000)
    ap.add_argument("--seed", type=int, default=15015)
    ap.add_argument("--data", default="")
    args = ap.parse_args()

    outdir = Path("results")
    outdir.mkdir(exist_ok=True)
    data_path = Path(args.data) if args.data else outdir / "external_parrotta_data.xlsx"
    if not data_path.exists():
        print(f"Downloading pinned public dataset: {DATA_URL}")
        urllib.request.urlretrieve(DATA_URL, data_path)

    rows, header = load_rows(data_path)
    print("Columns:", header)
    print("Rows loaded:", len(rows))
    print("Experiments:", sorted({r["Exp"] for r in rows}))
    print("Feedback levels:", sorted({r["Feedback"] for r in rows}))

    result = {
        "status": "EXPLORATORY_REANALYSIS_NOT_PREREGISTERED",
        "interpretation_firewall": (
            "This reanalysis tests whether measured anticipatory cardiac change statistically mediates "
            "the effect of manipulated cardiac feedback on pain ratings. HR is not assumed to be the "
            "complete bodily/evaluative source, mediation assumptions are not fully guaranteed, and "
            "no result establishes which physical state is phenomenal."
        ),
        "source": {
            "paper": "Parrotta et al. (2026), Exposure to false cardiac feedback alters pain perception and anticipatory cardiac frequency",
            "doi": "10.7554/eLife.90013",
            "repository": "EP171993/painperception",
            "dataset_commit": DATA_COMMIT,
            "dataset_blob_sha": "447fff57c91fb5ff761449d98d93737e26b1236c",
            "dataset_url": DATA_URL,
        },
        "method": {
            "contrast": "Faster (1) vs Slower (0) feedback",
            "mediator": "trial-level HR measure supplied in public dataset",
            "outcomes": ["LIK (unpleasantness)", "VAS (pain intensity)"],
            "covariates": ["stimulus intensity centered as in authors' code", "standardized trial index"],
            "subject_control": "subject fixed effects",
            "uncertainty": f"{args.bootstrap}-replicate cluster bootstrap resampling subjects",
            "estimand_warning": "a*b is exploratory statistical mediation, not a proven natural direct/indirect causal effect.",
        },
        "analyses": {},
    }

    for exp_i, exp in enumerate(("Intero", "Extero")):
        result["analyses"][exp] = {}
        for out_i, outcome in enumerate(("LIK", "VAS")):
            print(f"Analyzing {exp}/{outcome} ...")
            result["analyses"][exp][outcome] = analyze(
                rows, exp, outcome, args.bootstrap, args.seed + 100 * exp_i + out_i
            )

    # Coarse gate interpretation only. It is deliberately conservative.
    gate = []
    for outcome in ("LIK", "VAS"):
        a = result["analyses"]["Intero"][outcome]
        cp_ci = a["cluster_bootstrap_by_subject"]["c_prime_direct_proxy_to_outcome"]
        ind_ci = a["cluster_bootstrap_by_subject"]["indirect_a_times_b"]
        total_ci = a["cluster_bootstrap_by_subject"]["c_total_proxy_to_outcome"]
        def excludes_zero(ci):
            return ci is not None and (ci["ci95"][0] > 0 or ci["ci95"][1] < 0)
        gate.append({
            "outcome": outcome,
            "total_effect_detected": excludes_zero(total_ci),
            "measured_HR_indirect_detected": excludes_zero(ind_ci),
            "residual_direct_detected_after_HR": excludes_zero(cp_ci),
        })
    result["exploratory_gate"] = gate

    json_path = outdir / "human_proxy_mediation_gate.json"
    json_path.write_text(json.dumps(result, indent=2), encoding="utf-8")

    md = [
        "# Human proxy/source mediation gate",
        "",
        "**Exploratory, not preregistered. Not a consciousness detector.**",
        "",
        f"Dataset pinned to external commit `{DATA_COMMIT}`.",
        "",
    ]
    for exp in ("Intero", "Extero"):
        md += [f"## {exp}"]
        for outcome in ("LIK", "VAS"):
            a = result["analyses"][exp][outcome]
            p = a["point"]
            ci = a["cluster_bootstrap_by_subject"]
            md += [
                f"### {outcome}",
                f"- n = {p['n_rows']} trials, {p['n_subjects']} subjects",
                f"- total proxy effect c = {p['c_total_proxy_to_outcome']:.6g}",
                f"- proxy→HR a = {p['a_proxy_to_HR']:.6g}",
                f"- HR→outcome b = {p['b_HR_to_outcome_adjusted']:.6g}",
                f"- indirect a*b = {p['indirect_a_times_b']:.6g}",
                f"- residual/direct c' = {p['c_prime_direct_proxy_to_outcome']:.6g}",
                f"- mediated fraction = {p['mediated_fraction'] if p['mediated_fraction'] is not None else 'NA'}",
                f"- 95% bootstrap CI total = {ci['c_total_proxy_to_outcome']['ci95'] if ci['c_total_proxy_to_outcome'] else 'NA'}",
                f"- 95% bootstrap CI indirect = {ci['indirect_a_times_b']['ci95'] if ci['indirect_a_times_b'] else 'NA'}",
                f"- 95% bootstrap CI direct = {ci['c_prime_direct_proxy_to_outcome']['ci95'] if ci['c_prime_direct_proxy_to_outcome'] else 'NA'}",
                "",
            ]
    md += [
        "## Interpretation firewall",
        result["interpretation_firewall"],
        "",
        "A residual proxy coefficient after HR adjustment does not prove a source-independent phenomenal effect: the measured HR variable is only one physiological channel and standard mediation assumptions may fail.",
    ]
    (outdir / "human_proxy_mediation_gate.md").write_text("\n".join(md), encoding="utf-8")
    print(json.dumps(result["exploratory_gate"], indent=2))


if __name__ == "__main__":
    main()
